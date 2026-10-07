"""Bounded local extraction. Files never become executable web content."""
from io import BytesIO
from pathlib import Path
import warnings
import zipfile
from PIL import Image, ImageOps
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import pytesseract

MAX_TEXT = 180000
ALLOWED = {'.txt', '.md', '.csv', '.tsv', '.pdf', '.docx', '.xlsx', '.png', '.jpg', '.jpeg', '.webp'}
MIMES = {'.pdf': 'application/pdf', '.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.webp': 'image/webp',
         '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}
Image.MAX_IMAGE_PIXELS = 25_000_000


def validate_file(name, data):
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED: raise ValueError('Поддерживаются PDF, DOCX, XLSX, CSV, TXT, MD, PNG, JPG и WebP. Старые DOC/XLS сохраните в новом формате.')
    if not data: raise ValueError('Файл пустой')
    if ext in {'.docx', '.xlsx'}:
        with zipfile.ZipFile(BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 60*1024*1024 or len(z.infolist()) > 3000:
                raise ValueError('Слишком большой объём распакованного документа')
    if ext == '.pdf' and not data.lstrip().startswith(b'%PDF-'): raise ValueError('Содержимое не соответствует формату PDF')
    if ext in {'.png', '.jpg', '.jpeg', '.webp'}:
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(BytesIO(data)) as im: im.verify()
    return ext


def ocr(image):
    image = ImageOps.exif_transpose(image).convert('RGB')
    image.thumbnail((2800, 2800))
    langs = pytesseract.get_languages()
    lang = '+'.join(x for x in ['rus', 'eng'] if x in langs)
    if 'rus' not in langs: raise ValueError('На сервере нет русского OCR. Установите tesseract-ocr-rus или вставьте распознанный текст вручную.')
    return pytesseract.image_to_string(image, lang=lang, timeout=45)


def extract(name, data):
    ext = validate_file(name, data)
    note = ''
    if ext in {'.txt', '.md', '.csv', '.tsv'}:
        try: text = data.decode('utf-8-sig')
        except UnicodeDecodeError: text = data.decode('cp1251')
    elif ext == '.docx':
        doc = Document(BytesIO(data))
        chunks = [p.text for p in doc.paragraphs]
        for ti, table in enumerate(doc.tables, 1):
            chunks.append(f'Таблица {ti}')
            chunks.extend(' | '.join(c.text for c in r.cells) for r in table.rows)
        text = '\n'.join(chunks)
    elif ext == '.xlsx':
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        chunks = []
        total = 0
        try:
            for ws in wb.worksheets[:30]:
                chunks.append(f'Лист: {ws.title}')
                for row in ws.iter_rows(max_row=min(ws.max_row or 1, 10000), max_col=min(ws.max_column or 1, 100), values_only=True):
                    line = ' | '.join('' if v is None else str(v) for v in row).strip(' |')
                    if line: chunks.append(line); total += len(line)
                    if total > MAX_TEXT: break
                if total > MAX_TEXT: break
        finally: wb.close()
        text = '\n'.join(chunks)
    elif ext == '.pdf':
        pdf = PdfReader(BytesIO(data))
        if pdf.is_encrypted: raise ValueError('Снимите пароль с PDF и загрузите его снова')
        if len(pdf.pages) > 60: raise ValueError('Разделите PDF на файлы до 60 страниц')
        chunks = []
        scanned = 0
        renderer = None
        try:
            for i, page in enumerate(pdf.pages):
                t = page.extract_text() or ''
                if len(t.strip()) < 35:
                    scanned += 1
                    if scanned > 15: raise ValueError('Разделите сканированный PDF на части до 15 страниц')
                    if renderer is None:
                        import pymupdf
                        renderer = pymupdf.open(stream=data, filetype='pdf')
                    pix = renderer[i].get_pixmap(dpi=130)
                    t = ocr(Image.open(BytesIO(pix.tobytes('png'))))
                chunks.append(f'Страница {i+1}\n{t}')
                if sum(map(len, chunks)) > MAX_TEXT: break
            text = '\n\n'.join(chunks)
        finally:
            if renderer: renderer.close()
        if scanned: note = 'Использовано распознавание сканов. Проверьте цифры, объёмы и названия перед генерацией.'
    else:
        with Image.open(BytesIO(data)) as im: text = ocr(im)
        note = 'Текст распознан с фотографии. Проверьте цифры, объёмы и названия перед генерацией.'
    if len(text) > MAX_TEXT: note += ' Достигнут лимит 180 000 символов. Разделите файл для полного обучения.'
    text = text[:MAX_TEXT].replace('\x00', '').strip()
    if len(text) < 20: raise ValueError('Не удалось распознать достаточно текста. Вставьте текст вручную или загрузите более чёткий файл.')
    return text, note.strip()
