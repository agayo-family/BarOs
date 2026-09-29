from pathlib import Path
from io import BytesIO, StringIO
import csv
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

MAX_CHARS = 180_000

def _cut(text: str) -> str:
    return text[:MAX_CHARS]

def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext in {".txt", ".md", ".csv", ".tsv"}:
            return _cut(data.decode("utf-8", errors="replace"))
        if ext == ".pdf":
            reader = PdfReader(BytesIO(data))
            return _cut("\n\n".join((p.extract_text() or "") for p in reader.pages))
        if ext == ".docx":
            doc = Document(BytesIO(data))
            return _cut("\n".join(p.text for p in doc.paragraphs))
        if ext == ".xlsx":
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
            chunks = []
            for ws in wb.worksheets:
                chunks.append(f"# Лист: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row]
                    if any(vals): chunks.append("\t".join(vals))
            return _cut("\n".join(chunks))
    except Exception as e:
        return f"[Не удалось извлечь текст автоматически: {e}]"
    return "[Файл сохранён, но автоматическое извлечение текста для этого формата пока не поддерживается.]"
