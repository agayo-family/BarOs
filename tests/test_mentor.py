import asyncio
import importlib
import io
import json
import os
import tempfile
import unittest
from datetime import timedelta
from unittest.mock import patch

import httpx
from PIL import Image
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select, func

os.environ.setdefault('DATABASE_URL', 'sqlite:///'+tempfile.mktemp(prefix='baros-mentor-', suffix='.db'))
os.environ['BAROS_WORKER'] = '0'
from baros.db import Base, engine, SessionLocal
from baros.models import Organization, Course
from baros.security import hash_password
from baros.v2.models import Account, VenueSettings, CourseSettings, CourseRevision, Exam, Job, AIRequest, MentorTurn, SourceFile, LearningProfile, LessonRead, CourseProgram
from baros.v2.domain import dump, parse, now
from baros.v2 import ai, ai_providers, ai_vision
from baros.v2.extraction import extract

app = importlib.import_module('baros.v2.app').app
TEXT = 'Лимонад Солнечный: сироп персика 20 мл, сок лимона 10 мл, содовая 120 мл. Подавайте со льдом.'
QUESTION = {'prompt': 'Сколько сиропа персика в лимонаде Солнечный?', 'choices': ['20 мл', '10 мл', '30 мл', '40 мл'], 'correct_index': 0, 'explanation': 'В составе указан сироп персика 20 мл.', 'type': 'knowledge'}
QWEN_ENV = {'AI_PROVIDER': 'qwen', 'MENTOR_PROVIDER': 'qwen', 'QWEN_ENABLED': '1', 'QWEN_API_KEY': 'test-secret-not-real',
            'QWEN_BASE_URL': 'https://fixture.ap-southeast-1.maas.aliyuncs.com/compatible-mode/v1', 'AI_YANDEX_FALLBACK': '0'}


def vendor(result, tokens=100, code=200):
    return httpx.Response(code, json={'choices': [{'message': {'content': json.dumps(result, ensure_ascii=False)}, 'finish_reason': 'stop'}],
                                       'usage': {'prompt_tokens': tokens, 'completion_tokens': 100}})


class MentorTests(unittest.TestCase):
    def setUp(self):
        # All vendor posts below are mocked; don't instantiate the host's optional SOCKS transport.
        self.env = patch.dict(os.environ, {'MENTOR_PROVIDER': 'auto', 'AI_YANDEX_FALLBACK': '0', 'QWEN_ENABLED': '0', 'YANDEX_ENABLED': '0',
                                          **{key: '' for key in ['ALL_PROXY', 'all_proxy', 'HTTP_PROXY', 'http_proxy', 'HTTPS_PROXY', 'https_proxy']}})
        self.env.start(); Base.metadata.drop_all(engine)
        self.client = TestClient(app); self.client.__enter__()
        with SessionLocal() as db:
            org = Organization(name='Учебный бар'); other = Organization(name='Чужой бар'); db.add_all([org, other]); db.flush(); self.oid, self.other = org.id, other.id
            db.add_all([VenueSettings(organization_id=org.id, join_code='MENTOR1'), VenueSettings(organization_id=other.id, join_code='MENTOR2')])
            self.ids = {}
            for name, role, oid in [('worker', 'employee', org.id), ('peer', 'employee', org.id), ('manager', 'manager', org.id), ('foreign', 'manager', other.id), ('owner', 'owner', None)]:
                a = Account(name=name, login=name, password_hash=hash_password('mentor-pass'), organization_id=oid, role=role, positions_json='["bartender"]'); db.add(a); db.flush(); self.ids[name] = a.id
            c = Course(organization_id=org.id, title='Лимонады', published=True); db.add(c); db.flush(); self.cid = c.id
            data = {'title': c.title, 'description': 'Меню', 'positions': ['bartender'], 'passing_score': 80, 'quiz_size': 1, 'time_limit_minutes': 5, 'max_attempts': 0, 'deadline_days': 7, 'due_at': None, 'required': True, 'is_intro': False, 'lessons': [{'title': 'Состав', 'body': TEXT}], 'questions': [QUESTION]}
            rev = CourseRevision(course_id=c.id, version=1, payload=dump(data)); db.add(rev); db.flush(); self.rev = rev.id
            db.add(CourseSettings(course_id=c.id, version=1, positions_json='["bartender"]'))
            secret = Course(organization_id=other.id, title='PRIVATE MENU', published=True); db.add(secret); db.flush(); self.foreign_cid = secret.id
            db.commit()
        self.login('worker'); self.post('/api/growth/pet', {'pet_id': 'ember'})

    def tearDown(self):
        self.client.__exit__(None, None, None); self.env.stop()

    def login(self, name):
        self.client.cookies.clear(); self.assertEqual(self.client.post('/api/auth/login', json={'login': name, 'password': 'mentor-pass'}).status_code, 200)
        self.headers = {'x-csrf-token': self.client.get('/api/me').json()['csrf']}
        if name == 'owner': self.headers['x-organization-id'] = str(self.oid)

    def post(self, path, data): return self.client.post(path, json=data, headers=self.headers)
    def message(self, text='Объясни сироп персика', nonce='message-unique-0001', **ctx):
        return self.post('/api/mentor/messages', {'course_id': self.cid, 'message': text, 'nonce': nonce, **ctx})

    def test_local_pet_context_history_idempotence_and_no_paid_games(self):
        with patch.object(httpx.AsyncClient, 'post') as paid:
            r = self.message(); self.assertEqual(r.status_code, 200, r.text); self.assertIn('20 мл', r.json()['response']['reply']); self.assertEqual(r.json()['response']['mode'], 'local')
            self.assertEqual(self.message().json(), r.json()); self.assertEqual(self.message('Другой вопрос').status_code, 409)
            history = self.client.get('/api/mentor', params={'course_id': self.cid}).json(); self.assertEqual(history['pet']['name'], 'Искрик'); self.assertEqual(len(history['turns']), 1)
            self.client.get(f'/api/learning/{self.cid}/cards'); self.client.get('/api/games'); paid.assert_not_called()
        with SessionLocal() as db:
            self.assertEqual(db.scalar(select(func.count(AIRequest.id))), 0); self.assertEqual(db.get(LearningProfile, self.ids['worker']).xp, 0)

    def test_authorisation_csrf_roles_pet_swap_and_revision(self):
        self.message(); self.assertEqual(self.client.post('/api/mentor/messages', json={'message': 'Hi', 'nonce': 'missing-csrf-1'}).status_code, 403)
        self.assertEqual(self.message(preview_pet_id='moon', nonce='preview-invalid1').status_code, 403)
        self.assertEqual(self.client.get('/api/mentor', params={'course_id': self.foreign_cid}).status_code, 404)
        self.login('peer'); self.post('/api/growth/pet', {'pet_id': 'moon'}); self.assertEqual(self.client.get('/api/mentor', params={'course_id': self.cid}).json()['turns'], [])
        self.login('worker'); self.post('/api/growth/pet', {'pet_id': 'prism', 'confirm_swap': True}); self.assertEqual(self.client.get('/api/mentor', params={'course_id': self.cid}).json()['turns'], [])
        self.login('foreign'); self.assertEqual(self.client.get('/api/mentor', params={'course_id': self.cid}).status_code, 404); self.assertEqual(self.client.get('/api/ai/spend').status_code, 200)
        self.login('manager'); data = self.client.get('/api/mentor', params={'course_id': self.cid, 'preview_pet_id': 'sand'}).json(); self.assertTrue(data['preview']); self.assertEqual(data['pet']['name'], 'Зарри')

    def test_active_exam_block_and_own_finished_review(self):
        with SessionLocal() as db:
            db.add(Exam(id='mentor-exam', account_id=self.ids['worker'], revision_id=self.rev, questions_json=dump([QUESTION]), expires_at=now()+timedelta(minutes=5))); db.commit()
        self.assertEqual(self.client.get('/api/mentor').status_code, 409); self.assertEqual(self.message().status_code, 409)
        review = {**QUESTION, 'selected': 1, 'correct': False}
        with SessionLocal() as db:
            exam = db.get(Exam, 'mentor-exam'); exam.finished_at = now(); exam.result_json = dump({'passed': False, 'review': [review]}); db.commit()
        r = self.message(exam_id='mentor-exam', review_index=0); self.assertEqual(r.status_code, 200, r.text); self.assertIn('20 мл', r.json()['response']['reply'])
        self.login('peer'); self.post('/api/growth/pet', {'pet_id': 'ember'}); self.assertEqual(self.message(exam_id='mentor-exam', review_index=0).status_code, 404)

    def test_learning_memory_game_errors_and_context(self):
        with SessionLocal() as db:
            db.add(LessonRead(account_id=self.ids['worker'], revision_id=self.rev, lesson_index=0, completed_at=now())); db.commit()
        run = self.post(f'/api/games/{self.cid}/start', {'mode': 'mix', 'size': 5}).json(); q = run['questions'][0]
        # Mix may choose a variant; a scenario always uses choices for this knowledge bank.
        run = self.post(f'/api/games/{self.cid}/start', {'mode': 'words', 'size': 5}).json(); q = run['questions'][0]
        wrong = [x['id'] for x in q['tokens'] if x['text'] != '20'][:1]
        r = self.post('/api/games/runs/'+run['id']+'/answer', {'event_id': 'game-mistake', 'question_id': q['id'], 'value': wrong}); self.assertEqual(r.status_code, 200, r.text)
        data = self.client.get('/api/mentor', params={'course_id': self.cid}).json(); self.assertIn('Состав', data['memory']['learned']); self.assertGreater(data['memory']['weak'][0]['mistakes'], 0)
        r = self.message(game_id=run['id'], question_id=q['id']); self.assertEqual(r.status_code, 200, r.text); self.assertIn('20 мл', r.json()['response']['reply'])

    def test_daily_message_limits_and_progress_preservation(self):
        with patch.dict(os.environ, {'MENTOR_EMPLOYEE_DAILY_MESSAGES': '1', 'MENTOR_VENUE_DAILY_MESSAGES': '1'}):
            self.assertEqual(self.message().status_code, 200); self.assertEqual(self.message(nonce='message-next-0002').status_code, 429)
            self.login('peer'); self.post('/api/growth/pet', {'pet_id': 'ember'}); self.assertEqual(self.message(nonce='peer-message-0003').status_code, 429)
            self.assertEqual(self.client.get(f'/api/learning/{self.cid}/cards').status_code, 200)

    def test_qwen_payload_personality_cost_and_citations(self):
        async def answer(client, url, **kw):
            self.assertTrue(url.endswith('/compatible-mode/v1/chat/completions')); p = kw['json']; self.assertEqual(p['model'], 'qwen3.5-flash'); self.assertFalse(p['enable_thinking']); self.assertEqual(p['response_format']['type'], 'json_object')
            self.assertIn('Искрик', p['messages'][0]['content']); self.assertNotIn('PRIVATE MENU', dump(p))
            return vendor({'reply': 'В материале указано 20 мл сиропа персика.', 'kind': 'grounded', 'citations': [{'source_id': 'S1', 'quote': TEXT}], 'follow_up': 'Какую деталь проверим?'})
        with patch.dict(os.environ, QWEN_ENV), patch.object(httpx.AsyncClient, 'post', new=answer):
            r = self.message(); self.assertEqual(r.status_code, 200, r.text); self.assertEqual(r.json()['response']['mode'], 'ai'); self.message()
        with SessionLocal() as db:
            rows = db.scalars(select(AIRequest)).all(); self.assertEqual(len(rows), 1); self.assertEqual(rows[0].cost_micro, 50); self.assertEqual(rows[0].currency, 'USD')
            self.assertEqual(db.get(LearningProfile, self.ids['worker']).pet_id, 'ember')

    def test_no_automatic_activation_invalid_evidence_and_budget(self):
        with patch.dict(os.environ, {**QWEN_ENV, 'QWEN_ENABLED': '0'}), patch.object(httpx.AsyncClient, 'post') as call:
            self.message(); call.assert_not_called()
        async def bad(client, url, **kw): return vendor({'reply': 'Добавь 999 мл молока.', 'kind': 'grounded', 'citations': [{'source_id': 'S1', 'quote': 'Неподтверждённая цитата'}]})
        with patch.dict(os.environ, QWEN_ENV), patch.object(httpx.AsyncClient, 'post', new=bad):
            r = self.message(nonce='invalid-evidence1'); self.assertEqual(r.json()['response']['mode'], 'local_fallback'); self.assertNotIn('999', r.json()['response']['reply'])
        async def bad_follow_up(client,url,**kw):return vendor({'reply':'По материалу: сироп персика 20 мл.','kind':'grounded','citations':[{'source_id':'S1','quote':TEXT}],'follow_up':'Добавим 999 мл молока?'})
        with patch.dict(os.environ,QWEN_ENV),patch.object(httpx.AsyncClient,'post',new=bad_follow_up):
            r=self.message(nonce='invalid-followup1');self.assertEqual(r.json()['response']['mode'],'local_fallback');self.assertNotIn('999',dump(r.json()['response']))
        with patch.dict(os.environ, {**QWEN_ENV, 'AI_EMPLOYEE_DAILY_BUDGET_USD': '0'}), patch.object(httpx.AsyncClient, 'post') as call:
            r = self.message(nonce='budget-blocked-1'); self.assertEqual(r.json()['response']['mode'], 'local_fallback'); call.assert_not_called()

    def test_yandex_fallback_native_currency_and_no_auth_retry(self):
        calls = []
        async def flaky(client, url, **kw):
            calls.append((url, kw));
            if len(calls) == 1: raise httpx.ReadTimeout('Test timeout')
            return vendor({'reply': 'В материале указан сироп персика 20 мл.', 'kind': 'grounded', 'citations': [{'source_id': 'S1', 'quote': TEXT}]})
        env = {**QWEN_ENV, 'AI_YANDEX_FALLBACK': '1', 'YANDEX_ENABLED': '1', 'YANDEX_API_KEY': 'test-yandex-secret', 'YANDEX_FOLDER_ID': 'test-folder'}
        with patch.dict(os.environ, env), patch.object(httpx.AsyncClient, 'post', new=flaky):
            r = self.message(); self.assertEqual(r.status_code, 200, r.text); self.assertEqual(r.json()['response']['provider'], 'yandex')
        self.assertEqual(calls[1][1]['headers']['Authorization'], 'Api-Key test-yandex-secret'); self.assertEqual(calls[1][1]['headers']['OpenAI-Project'], 'test-folder')
        with SessionLocal() as db:
            rows = db.scalars(select(AIRequest).order_by(AIRequest.created_at)).all(); self.assertEqual([r.currency for r in rows], ['USD', 'RUB']); self.assertEqual(rows[0].status, 'unconfirmed'); self.assertEqual(rows[1].cost_micro, 30000)
        async def rejected(client, url, **kw): return httpx.Response(401, json={})
        with patch.dict(os.environ, env), patch.object(httpx.AsyncClient, 'post', new=rejected): self.message(nonce='rejected-auth-001')
        with SessionLocal() as db: self.assertEqual(db.scalar(select(func.count(AIRequest.id))), 3)

    def test_documents_photo_checkpoint_draft_and_offline_practice(self):
        doc = Document(); doc.add_paragraph(TEXT); buf = io.BytesIO(); doc.save(buf); self.assertIn('20 мл', extract('menu.docx', buf.getvalue())[0])
        image = Image.new('RGB', (2400, 1800), 'white'); raw = io.BytesIO(); image.save(raw, 'JPEG')
        images, _ = ai_vision.pages('menu.jpg', raw.getvalue()); decoded = images[0][1].split(',', 1)[1]
        import base64
        clean = Image.open(io.BytesIO(base64.b64decode(decoded))); self.assertLessEqual(max(clean.size), 1600)
        doc.add_picture(io.BytesIO(raw.getvalue())); photo_doc=io.BytesIO();doc.save(photo_doc)
        mixed=SourceFile(name='mixed.docx',content=photo_doc.getvalue(),text=TEXT,warning='',extraction_status='ready')
        self.assertTrue(ai_vision.needs_vision(mixed));self.assertEqual(len(ai_vision.pages(mixed.name,mixed.content)[0]),1)
        mixed.warning='Текст проверен и отредактирован вручную';self.assertFalse(ai_vision.needs_vision(mixed))
        with SessionLocal() as db:
            src = SourceFile(organization_id=self.oid, name='menu.jpg', mime='image/jpeg', size=len(raw.getvalue()), content=raw.getvalue(), text='', extraction_status='needs_text'); db.add(src); db.flush(); sid = src.id
            j = Job(id='vision-job', organization_id=self.oid, account_id=self.ids['manager'], context='Тема: меню', positions_json='["bartender"]', source_ids_json=dump([sid]), checkpoint=dump({'vision_ids': [sid]}), lease_token='lease'); db.add(j); db.commit()
        counter = {'vision': 0}
        async def model(client, url, **kw):
            p = kw['json']; self.assertEqual(p['model'], 'qwen3.5-plus'); content = p['messages'][-1]['content']
            if isinstance(content, list):
                counter['vision'] += 1; self.assertTrue(content[1]['image_url']['url'].startswith('data:image/jpeg;base64,')); return vendor({'text': TEXT, 'uncertain': []})
            if content.startswith('Создай'): return vendor({'title': 'Лимонады', 'description': 'Учимся меню', 'lessons': [{'title': 'Состав', 'body': TEXT}], 'limitations': []})
            return vendor({'questions': [{**QUESTION, 'source_quote': TEXT, 'lesson_index': 0}], 'limitations': ['Недостаточно разных фактов']})
        with patch.dict(os.environ, QWEN_ENV), patch.object(ai, 'AI_PROVIDER', 'qwen'), patch.object(httpx.AsyncClient, 'post', new=model): asyncio.run(ai.run_job('vision-job', 'lease'))
        with SessionLocal() as db:
            j = db.get(Job, 'vision-job'); self.assertEqual(j.status, 'needs_review'); self.assertEqual(parse(j.checkpoint)['vision_done'], [sid]); self.assertIn('20 мл', db.get(SourceFile, sid).text); self.assertIsNotNone(db.get(CourseProgram, j.course_id)); self.assertFalse(db.get(Course, j.course_id).published)
            self.assertEqual(counter['vision'], 1); self.assertTrue(all(r.cost_micro is not None for r in db.scalars(select(AIRequest)).all()))

    def test_spend_visibility_unknown_prices_and_currency_guards(self):
        self.assertEqual(self.client.get('/api/ai/spend').status_code, 403)
        with patch.dict(os.environ, {**QWEN_ENV, 'QWEN_CHAT_MODEL': 'unknown'}): self.assertFalse(ai_providers.public_status()['configured'])
        self.login('manager'); usage = self.client.get('/api/ai/spend').json(); self.assertEqual(usage['currencies']['USD']['daily'], 0)
        self.assertNotIn('test-secret', dump(usage))
        with patch.dict(os.environ, {'AI_DAILY_BUDGET_USD': 'NaN'}):
            with self.assertRaises(ValueError): ai_providers.budgets('USD')

    def test_interrupted_message_recovery_and_invalid_context(self):
        first=self.message()
        with SessionLocal() as db:
            turn=db.get(MentorTurn,first.json()['id']);turn.status='pending';turn.response_json='{}';turn.created_at=now()-timedelta(minutes=5);db.commit()
        with patch.object(httpx.AsyncClient,'post') as paid:
            data=self.client.get('/api/mentor',params={'course_id':self.cid}).json()
            self.assertEqual(data['turns'][0]['status'],'interrupted')
            self.assertEqual(self.message().json()['status'],'interrupted');paid.assert_not_called()
        self.assertEqual(self.client.get('/api/mentor',params={'course_id':-1}).status_code,422)
        self.assertEqual(self.client.get('/api/mentor',params={'lesson_index':-1}).status_code,422)

    def test_revision_change_hides_old_memory_and_provider_config_error(self):
        self.message()
        with SessionLocal() as db:
            db.add(LessonRead(account_id=self.ids['worker'],revision_id=self.rev,lesson_index=0,completed_at=now()))
            data=parse(db.get(CourseRevision,self.rev).payload);data['lessons'][0]['title']='Новая редакция'
            db.add(CourseRevision(course_id=self.cid,version=2,payload=dump(data)));db.get(CourseSettings,self.cid).version=2;db.commit()
        data=self.client.get('/api/mentor',params={'course_id':self.cid}).json()
        self.assertEqual(data['turns'],[]);self.assertEqual(data['memory']['learned'],[])
        self.login('manager')
        with patch.dict(os.environ,{**QWEN_ENV,'QWEN_BASE_URL':'https://untrusted.example/v1'}),patch.object(ai,'AI_PROVIDER','qwen'):
            r=self.post('/api/jobs',{'title':'Тема','positions':['bartender'],'source_ids':[],'interview':True})
            self.assertEqual(r.status_code,503,r.text)


if __name__ == '__main__': unittest.main()
