import asyncio
import importlib
import os
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

os.environ.setdefault('DATABASE_URL','sqlite:///'+tempfile.mktemp(prefix='baros-experience-',suffix='.db'))
os.environ['BAROS_WORKER']='0'

import httpx
from fastapi.testclient import TestClient
from sqlalchemy import select,func
from baros.db import Base,engine,SessionLocal
from baros.models import Organization,Course,Question,Lesson
from baros.security import hash_password
from baros.v2 import ai
from baros.v2.content import job_card
from baros.v2.domain import dump,parse
from baros.v2.models import Account,VenueSettings,CourseSettings,CourseRevision,Exam,LessonRead,Job

module=importlib.import_module('baros.v2.app')
QUOTE='Чай заваривается при температуре 90 °C.'
QUESTION={'prompt':'Какую температуру воды используют для заваривания чая по стандарту?',
          'choices':['90 °C','80 °C','70 °C','100 °C'],'correct_index':0,
          'explanation':'В материале задана температура воды 90 °C.','type':'knowledge',
          'source_quote':QUOTE,'lesson_index':0}
FIRST={'prompt':'Как называется напиток с заданной температурой заваривания?',
       'choices':['Чай','Кофе','Какао','Матча'],'correct_index':0,
       'explanation':'Температура 90 °C указана для чая.','type':'knowledge',
       'source_quote':QUOTE,'lesson_index':0}


class LearningExperienceTests(unittest.TestCase):
    def setUp(self):
        Base.metadata.drop_all(engine)
        self.client=TestClient(module.app);self.client.__enter__()
        with SessionLocal() as db:
            org=Organization(name='Our venue');other=Organization(name='Another venue');db.add_all([org,other]);db.flush()
            db.add_all([VenueSettings(organization_id=org.id,join_code='OURCODE123'),VenueSettings(organization_id=other.id,join_code='OTHERCODE12')])
            a=Account(name='Learner',login='learner',role='employee',organization_id=org.id,
                      password_hash=hash_password('test-only-password'),positions_json='["bartender"]')
            course=Course(organization_id=org.id,title='Draft title',published=True)
            foreign=Course(organization_id=other.id,title='Foreign',published=True)
            db.add_all([a,course,foreign]);db.flush()
            db.add(CourseSettings(course_id=course.id,version=1))
            payload={'title':'Published title','description':'Published description','positions':['bartender'],
                     'questions':[QUESTION],'lessons':[{'title':'Theory','body':QUOTE}],
                     'passing_score':80,'required':True,'quiz_size':1,'time_limit_minutes':10,
                     'max_attempts':0,'deadline_days':7,'due_at':None,'is_intro':False}
            db.add(CourseRevision(course_id=course.id,version=1,payload=dump(payload)))
            db.add(Question(course_id=course.id,prompt='Unpublished draft question',choices_json='["Private answer","a","b","c"]',correct_index=0))
            db.commit();self.cid,self.foreign,self.aid,self.oid=course.id,foreign.id,a.id,org.id
        self.assertEqual(self.client.post('/api/auth/login',json={'login':'learner','password':'test-only-password'}).status_code,200)

    def tearDown(self):
        self.client.__exit__(None,None,None)

    def test_cards_use_published_revision_without_exam_or_read_credit(self):
        response=self.client.get(f'/api/learning/{self.cid}/cards');self.assertEqual(response.status_code,200)
        data=response.json();self.assertEqual(data['title'],'Published title')
        self.assertEqual(data['cards'][0]['answer'],'90 °C')
        self.assertNotIn('Private answer',response.text)
        self.assertEqual(self.client.get('/api/learning').json()['courses'][0]['card_count'],1)
        with SessionLocal() as db:
            self.assertEqual(db.scalar(select(func.count(Exam.id))),0)
            self.assertEqual(db.scalar(select(func.count(LessonRead.id))),0)

    def test_cards_respect_venue_position_archive_and_subscription(self):
        self.assertEqual(self.client.get(f'/api/learning/{self.foreign}/cards').status_code,404)
        with SessionLocal() as db:
            a=db.get(Account,self.aid);a.positions_json='["waiter"]';db.commit()
        self.assertEqual(self.client.get(f'/api/learning/{self.cid}/cards').status_code,403)
        with SessionLocal() as db:
            db.get(Account,self.aid).positions_json='["bartender"]';db.get(CourseSettings,self.cid).archived=True;db.commit()
        self.assertEqual(self.client.get(f'/api/learning/{self.cid}/cards').status_code,404)
        with SessionLocal() as db:
            db.get(CourseSettings,self.cid).archived=False;db.get(VenueSettings,self.oid).subscription_status='past_due';db.commit()
        self.assertEqual(self.client.get(f'/api/learning/{self.cid}/cards').status_code,402)

    def job(self):
        with SessionLocal() as db:
            j=Job(id='resume-test',organization_id=self.oid,account_id=self.aid,context=QUOTE,
                  target=2,positions_json='["bartender"]',lease_token='test-lease',status='running')
            cp={'title':'AI draft','lessons':[{'title':'Tea','body':QUOTE}],'questions':[FIRST]}
            j.checkpoint=dump(cp);db.add(j);db.flush();ai.persist_course(db,j,cp);db.commit();return j.course_id

    def test_resuming_ai_updates_existing_draft_and_completes_real_bank(self):
        cid=self.job()
        with patch.object(ai,'call_model',new=AsyncMock(return_value={'questions':[QUESTION],'_model_used':ai.FREE_BACKUP})):
            asyncio.run(ai.run_job('resume-test','test-lease'))
        with SessionLocal() as db:
            j=db.get(Job,'resume-test');self.assertEqual(j.course_id,cid);self.assertEqual(j.status,'succeeded')
            self.assertEqual(db.scalar(select(func.count(Question.id)).where(Question.course_id==cid)),2)
            self.assertEqual(db.scalar(select(func.count(Lesson.id)).where(Lesson.course_id==cid)),1)
            card=job_card(j);self.assertEqual(card['progress'],100);self.assertEqual(card['questions_progress'],100)
            self.assertEqual(card['model_used'],ai.FREE_BACKUP)

    def test_incomplete_bank_and_edited_draft_are_not_overwritten(self):
        cid=self.job()
        with patch.object(ai,'call_model',new=AsyncMock(return_value={'questions':[],'limitations':['Недостаточно фактов']})):
            asyncio.run(ai.run_job('resume-test','test-lease'))
        with SessionLocal() as db:
            j=db.get(Job,'resume-test');self.assertEqual(j.status,'needs_review')
            self.assertEqual(job_card(j)['questions_progress'],50);self.assertLess(job_card(j)['progress'],100)
            db.get(CourseSettings,cid).edit_version=2;j.lease_token='test-lease';j.status='running';db.commit()
        with patch.object(ai,'call_model',new=AsyncMock(return_value={'questions':[QUESTION]})):
            asyncio.run(ai.run_job('resume-test','test-lease'))
        with SessionLocal() as db:
            self.assertEqual(db.get(Job,'resume-test').status,'failed')
            self.assertEqual(db.scalar(select(func.count(Question.id)).where(Question.course_id==cid)),1)

    def test_free_routing_and_response_validation(self):
        self.job();sent=[];client_class=httpx.AsyncClient
        def response(request):
            import json
            sent.append(json.loads(request.content))
            return httpx.Response(200,json={'model':ai.FREE_BACKUP,'choices':[{'message':{'content':'{"lessons":[]}'}}]})
        with patch.object(ai,'OPENROUTER_API_KEY','unit-test-key'),patch.object(ai,'OPENROUTER_MODEL','openrouter/free'),patch.object(ai.httpx,'AsyncClient',side_effect=lambda **kw:client_class(transport=httpx.MockTransport(response),**kw)):
            result=asyncio.run(ai.call_model('Test only','resume-test','test-lease'))
        self.assertEqual(result['_model_used'],ai.FREE_BACKUP)
        self.assertEqual(sent[0]['models'],[ai.FREE_PRIMARY,ai.FREE_BACKUP,'openrouter/free'])
        self.assertTrue(sent[0]['provider']['require_parameters'])
        with patch.object(ai,'OPENROUTER_API_KEY','unit-test-key'),patch.object(ai,'OPENROUTER_MODEL','paid/model'),patch.object(ai,'AI_ALLOW_PAID_FALLBACK',False):
            self.assertIsNone(ai.provider())
        with self.assertRaisesRegex(ValueError,'некорректный JSON'):ai.decode_json('not JSON')
        self.assertEqual(ai.decode_json('```json\n{"ok":true}\n```'),{'ok':True})

    def test_upload_formats_ocr_fallback_and_ai_job_creation(self):
        from io import BytesIO
        from docx import Document
        from openpyxl import Workbook
        from PIL import Image
        import pymupdf
        with SessionLocal() as db:
            db.get(Account,self.aid).role='manager';db.commit()
        headers={'x-csrf-token':self.client.get('/api/me').json()['csrf']}
        doc=Document();doc.add_paragraph(QUOTE);word=BytesIO();doc.save(word)
        workbook=Workbook();workbook.active.append(['Чай','90 °C',QUOTE]);sheet=BytesIO();workbook.save(sheet);workbook.close()
        pdf=pymupdf.open();pdf.new_page().insert_text((30,40),'House tea is brewed at 90 degrees according to the venue standard.');pdf_bytes=pdf.tobytes();pdf.close()
        samples={'menu.txt':QUOTE.encode(),'menu.csv':('Напиток;Температура\nЧай;90 °C\n'+QUOTE).encode(),
                 'menu.docx':word.getvalue(),'menu.xlsx':sheet.getvalue(),'menu.pdf':pdf_bytes}
        ids=[]
        for name,raw in samples.items():
            r=self.client.post('/api/sources',files={'file':(name,raw)},data={'roles':'["bartender"]'},headers=headers)
            self.assertEqual(r.status_code,200,r.text);self.assertEqual(r.json()['status'],'ready');ids.append(r.json()['id'])
            self.assertIn('90',self.client.get('/api/sources/'+str(ids[-1])).json()['text'])
        self.assertEqual(self.client.post('/api/sources',files={'file':('bad.exe',b'bad')},data={'roles':'["bartender"]'},headers=headers).status_code,400)
        picture=BytesIO();Image.new('RGB',(20,20),'white').save(picture,format='PNG')
        with patch('baros.v2.extraction.ocr',side_effect=ValueError('Русский OCR недоступен. Вставьте текст вручную.')):
            photo=self.client.post('/api/sources',files={'file':('menu.png',picture.getvalue())},data={'roles':'["bartender"]'},headers=headers)
        self.assertEqual(photo.status_code,200);self.assertEqual(photo.json()['status'],'needs_text')
        sid=photo.json()['id'];self.assertIn('вручную',photo.json()['warning'])
        fixed=self.client.put('/api/sources/'+str(sid),json={'text':QUOTE,'positions':['bartender']},headers=headers)
        self.assertEqual(fixed.status_code,200);self.assertEqual(fixed.json()['status'],'ready')
        with patch.object(ai,'provider',return_value={'free':True}):
            generated=self.client.post('/api/jobs',json={'title':'Tea training','source_ids':[sid],'positions':['bartender'],'interview':False},headers=headers)
        self.assertEqual(generated.status_code,200,generated.text);self.assertEqual(generated.json()['target'],100)
        self.assertEqual(generated.json()['progress'],0)


if __name__=='__main__':unittest.main()
